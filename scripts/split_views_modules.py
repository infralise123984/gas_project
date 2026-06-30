"""Splitter: views_monolith_backup.py -> modular views/services (with decorators)."""
from __future__ import annotations

import ast
import textwrap
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "mockups" / "views_monolith_backup.py"

MODULE_MAP: dict[str, list[str]] = {
    "mockups/views/auth.py": [
        "index", "login_view", "logout_view", "perfil_view",
        "verificar_2fa_view", "activar_2fa_view", "desactivar_2fa_view", "crear_usuario",
    ],
    "mockups/views/catalogos.py": [
        "historial_precios", "gestionar_balones_lista", "gestionar_balones_crear",
        "gestionar_balones_editar", "gestionar_balones_eliminar", "auditoria_lista",
        "gestionar_sectores_lista", "gestionar_sectores_crear", "gestionar_sectores_editar",
        "gestionar_sectores_eliminar",
    ],
    "mockups/views/pedidos.py": [
        "transaccional_pedido", "editar_pedido", "mis_pedidos_hoy", "mis_pedidos_hoy_api",
        "consultas_pedidos", "detalle_pedido", "telefonista_cancelar_pedido",
    ],
    "mockups/views/entregas.py": [
        "mis_entregas_camionero", "mis_entregas_camionero_api", "camionero_historial",
        "camionero_historial_dia", "camionero_entregas", "camionero_entregas_api",
        "camionero_tomar_pedido", "camionero_marcar_entregado", "camionero_cancelar_entrega",
        "tarreo_pedido",
    ],
    "mockups/views/reportes.py": ["reporte_ventas", "reporte_sobres"],
    "mockups/views/sobres.py": [
        "lista_sobres_diarios", "editar_sobre_diario", "refrescar_sobre_diario",
        "crear_sobre_nuevo", "crear_sobre_post_cierre", "historial_sobres",
        "imprimir_sobre_diario", "exportar_sobre_excel",
    ],
    "mockups/services/exports.py": ["exportar_pedidos_excel"],
    "mockups/views/push.py": [
        "push_subscribe", "push_unsubscribe", "push_test", "push_status", "service_worker",
    ],
}

VIEW_HEADERS = textwrap.dedent('''\
    """Vistas HTTP — {doc}."""

    import io
    import logging
    import uuid
    from calendar import monthrange
    from datetime import date, datetime, time, timedelta
    from json import dumps
    from zoneinfo import ZoneInfo

    import openpyxl
    import pyotp
    import segno
    from django.contrib import messages
    from django.contrib.auth import authenticate, login, logout
    from django.contrib.auth.decorators import login_required
    from django.core.paginator import Paginator
    from django.db import transaction
    from django.db.models import Count, F, Q, Sum, Case, When, Value, IntegerField
    from django.http import HttpResponse, JsonResponse
    from django.shortcuts import get_object_or_404, redirect, render
    from django.urls import reverse
    from django.utils import timezone
    from django.utils.http import url_has_allowed_host_and_scheme
    from django.views.decorators.cache import never_cache
    from django.views.decorators.http import require_POST
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    from mockups.forms import (
        Activar2FAConfirmForm,
        CambiarPasswordForm,
        CrearUsuarioSeguroForm,
        Desactivar2FAForm,
        DetalleFormSet,
        DetalleFormSetEdit,
        DetallePedidoForm,
        EditarPerfilForm,
        LineaGastoFormSet,
        LineaPagoFormSet,
        LineaSobreFormSet,
        PedidoCabeceraForm,
        SectorForm,
        TipoBalonForm,
        Verificar2FAForm,
    )
    from mockups.models import (
        AuditoriaAccion,
        DetallePedido,
        HistorialCambioPedido,
        HistorialEstadoPedido,
        HistorialPrecioBalon,
        LineaSobre,
        Pedido,
        Sector,
        SobreDiario,
        TipoBalon,
        Usuario,
    )
    from mockups.services.camionero import (
        _kilos_de_pedido,
        _stats_dia_camionero,
        filtro_actividad_camionero_q,
        queryset_actividad_camionero_dia,
        stats_ventas_camionero,
    )
    from mockups.services.catalogos import get_balones_activos_ordenados
    {exports_import}from mockups.services.sobres import get_pedidos_queryset_para_sobre, sincronizar_sobre_desde_pedidos
    from mockups.utils.fechas import (
        MESES_ES_CAMIONERO,
        get_rango_utc_para_fecha,
        navegacion_mes,
        now_chile,
        parse_fecha_rango,
        parse_mes_param,
        rango_dia_chile,
        today_chile,
    )
    from mockups.utils.permisos import get_client_ip, get_display_name, require_roles

    security_logger = logging.getLogger('security')
    audit_logger = logging.getLogger('audit')


''')

EXPORTS_HEADER = textwrap.dedent('''\
    """Generación de exportaciones Excel."""

    from datetime import datetime

    import openpyxl
    from django.contrib import messages
    from django.db.models import Case, IntegerField, Sum, Value, When
    from django.http import HttpResponse
    from django.shortcuts import get_object_or_404, redirect
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    from mockups.models import Pedido, SobreDiario
    from mockups.services.catalogos import get_balones_activos_ordenados
    from mockups.utils.fechas import today_chile
    from mockups.utils.permisos import get_display_name


''')


def extract_function_source(source: str, node: ast.FunctionDef) -> str:
    segment = ast.get_source_segment(source, node)
    if not segment:
        raise ValueError(f"Could not extract source for {node.name}")
    # Include decorators and preceding comment lines
    start = node.lineno
    while start > 1:
        prev = source.splitlines()[start - 2]
        stripped = prev.strip()
        if stripped.startswith("@") or stripped == "" or stripped.startswith("#"):
            start -= 1
        else:
            break
    prefix = "\n".join(source.splitlines()[start - 1 : node.lineno - 1])
    if prefix.strip():
        return prefix + "\n" + segment + "\n\n"
    return segment + "\n\n"


def header_for(module: str) -> str:
    if module.endswith("exports.py"):
        return EXPORTS_HEADER
    doc = Path(module).stem
    exports_import = ""
    if module in ("mockups/views/pedidos.py", "mockups/views/reportes.py", "mockups/views/sobres.py"):
        exports_import = "from mockups.services.exports import exportar_pedidos_excel\n"
    camionero = (
        "from mockups.services.camionero import (\n"
        "    _kilos_de_pedido,\n"
        "    _stats_dia_camionero,\n"
        "    filtro_actividad_camionero_q,\n"
        "    queryset_actividad_camionero_dia,\n"
        "    stats_ventas_camionero,\n"
        ")\n"
    )
    sobres_import = "from mockups.services.sobres import get_pedidos_queryset_para_sobre, sincronizar_sobre_desde_pedidos\n"
    body = VIEW_HEADERS.format(doc=doc, exports_import=exports_import)
    if module == "mockups/views/auth.py":
        body = body.replace(camionero, "")
        body = body.replace(sobres_import, "")
    elif module == "mockups/views/push.py":
        body = body.replace(
            "from mockups.services.exports import exportar_pedidos_excel\n", ""
        )
        body = body.replace(camionero, "")
        body = body.replace(sobres_import, "")
    elif module == "mockups/views/catalogos.py":
        body = body.replace(camionero, "")
        body = body.replace(sobres_import, "")
    elif module == "mockups/views/sobres.py":
        body = body.replace(camionero, "")
    elif module == "mockups/views/pedidos.py":
        body = body.replace(camionero, "")
        body = body.replace(sobres_import, "")
    elif module == "mockups/views/entregas.py":
        body = body.replace(sobres_import, "")
    elif module == "mockups/views/reportes.py":
        body = body.replace(camionero, "")
        body = body.replace(sobres_import, "")
    return body


def main() -> None:
    source = SOURCE.read_text(encoding="utf-8")
    tree = ast.parse(source)

    func_nodes: dict[str, ast.FunctionDef] = {}
    for node in tree.body:
        if isinstance(node, ast.FunctionDef):
            func_nodes[node.name] = node

    name_to_module: dict[str, str] = {}
    for module, names in MODULE_MAP.items():
        for name in names:
            name_to_module[name] = module

    missing = set(name_to_module) - set(func_nodes)
    if missing:
        raise SystemExit(f"Functions missing from backup: {sorted(missing)}")

    buckets: dict[str, list[str]] = {m: [] for m in MODULE_MAP}
    for name in sorted(name_to_module, key=lambda n: func_nodes[n].lineno):
        module = name_to_module[name]
        buckets[module].append(extract_function_source(source, func_nodes[name]))

    for module, chunks in buckets.items():
        path = ROOT / module
        path.write_text(header_for(module) + "".join(chunks), encoding="utf-8")
        print(f"Wrote {path} ({len(chunks)} functions)")


if __name__ == "__main__":
    main()