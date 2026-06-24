"""Generación de exportaciones Excel (lógica sin request HTTP)."""

import openpyxl
from django.http import HttpResponse
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from mockups.utils.fechas import now_chile
from mockups.utils.permisos import get_display_name


def exportar_pedidos_excel(queryset, rango_fechas):
    """Genera archivo .xlsx con detalles de pedidos filtrados."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pedidos"

    # Configuración de estilos
    header_fill = PatternFill(start_color="0066CC", end_color="0066CC", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    # Títulos
    ws.merge_cells('A1:L1')
    ws['A1'] = f"Reporte de Pedidos - KIM GAS Rancagua"
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')

    if rango_fechas:
        ws.merge_cells('A2:L2')
        ws['A2'] = f"Período: {rango_fechas}"
        ws['A2'].alignment = Alignment(horizontal='center')

    # Cabeceras de tabla
    headers = ['ID', 'Fecha', 'Hora', 'Estado', 'Origen', 'Productos', 'Cantidad Total', 'Monto', 'Ganancia', 'Registrado por', 'Entregado por', 'Sector/Dirección']
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border

    # Llenado de datos (un solo recorrido; acumula totales en el mismo loop)
    row_num = 5
    total_monto = 0.0
    total_ganancia = 0.0
    for pedido in queryset:
        detalles = list(pedido.detalles.all())
        productos_list = [f"{d.cantidad}×{d.balon.nombre}" for d in detalles]
        cantidad_total = sum(d.cantidad for d in detalles)
        productos_str = " + ".join(productos_list) if productos_list else "—"
        ubicacion = f"{pedido.sector or '—'} / {pedido.direccion_entrega or '—'}"
        monto = float(pedido.monto_total)
        ganancia = float(pedido.ganancia_total)
        total_monto += monto
        total_ganancia += ganancia

        data = [
            pedido.id,
            pedido.fecha.strftime("%d/%m/%Y"),
            pedido.fecha.strftime("%H:%M"),
            pedido.get_estado_display(),
            pedido.get_origen_display(),
            productos_str,
            cantidad_total,
            monto,
            ganancia,
            get_display_name(pedido.registrador),
            get_display_name(pedido.entregador),
            ubicacion
        ]

        for col_num, value in enumerate(data, 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.value = value
            cell.border = border
            if col_num in [8, 9]:
                cell.number_format = '"$"#,##0'
            cell.alignment = Alignment(horizontal='center') if col_num in [1, 2, 3, 7, 8, 9] else Alignment(horizontal='left')
        row_num += 1

    # Totales finales
    ws.cell(row=row_num, column=1).value = "TOTALES:"
    ws.cell(row=row_num, column=1).font = Font(bold=True)

    ws.cell(row=row_num, column=8).value = total_monto
    ws.cell(row=row_num, column=8).number_format = '"$"#,##0'
    ws.cell(row=row_num, column=8).font = Font(bold=True)

    ws.cell(row=row_num, column=9).value = total_ganancia
    ws.cell(row=row_num, column=9).number_format = '"$"#,##0'
    ws.cell(row=row_num, column=9).font = Font(bold=True)

    # Ajuste de columnas
    column_widths = {'A': 8, 'B': 12, 'C': 10, 'D': 12, 'E': 18, 'F': 25, 'G': 12, 'H': 12, 'I': 12, 'J': 20, 'K': 20, 'L': 35}
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    timestamp = now_chile().strftime('%Y%m%d_%H%M')
    filename = f"pedidos_{timestamp}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response